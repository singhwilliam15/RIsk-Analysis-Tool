import numpy as np
import pandas as pd
import pytest

from var_calculator import (
    backtest_all_methods,
    binomial_cdf,
    calculate_all_var,
    calculate_cornish_fisher_var,
    calculate_ewma_var,
    calculate_historical_var,
    calculate_monte_carlo_var,
    calculate_parametric_var,
    calculate_student_t_var,
    christoffersen_test,
    estimate_beta,
    ewma_volatility,
    historical_worst_losses,
    perform_kupiec_backtest,
    rolling_historical_var,
    rolling_var_forecasts,
    run_stress_testing,
)


@pytest.fixture
def normal_returns():
    rng = np.random.default_rng(7)
    return pd.Series(rng.normal(0.0005, 0.015, 2000))


def test_parametric_var_matches_closed_form(normal_returns):
    res = calculate_parametric_var(normal_returns, 1_000_000, 0.99)
    mu, sigma = normal_returns.mean(), normal_returns.std(ddof=1)
    assert res["var_daily_pct"] == pytest.approx(2.326348 * sigma - mu, rel=1e-5)


def test_expected_shortfall_exceeds_var(normal_returns):
    for fn in (calculate_historical_var, calculate_parametric_var, calculate_monte_carlo_var):
        res = fn(normal_returns, 1_000_000, 0.95)
        assert res["cvar_daily_pct"] > res["var_daily_pct"] > 0


def test_methods_agree_on_normal_data(normal_returns):
    hist = calculate_historical_var(normal_returns, 1, 0.95)["var_daily_pct"]
    param = calculate_parametric_var(normal_returns, 1, 0.95)["var_daily_pct"]
    mc = calculate_monte_carlo_var(normal_returns, 1, 0.95, num_simulations=50_000, seed=1)["var_daily_pct"]
    assert hist == pytest.approx(param, rel=0.1)
    assert mc == pytest.approx(param, rel=0.05)


def test_holding_period_uses_square_root_of_time(normal_returns):
    one_day = calculate_historical_var(normal_returns, 1_000_000, 0.95, 1)
    ten_day = calculate_historical_var(normal_returns, 1_000_000, 0.95, 10)
    assert ten_day["var_scaled_amount"] == pytest.approx(one_day["var_daily_amount"] * np.sqrt(10))


def test_rolling_var_is_out_of_sample(normal_returns):
    rolling = rolling_historical_var(normal_returns, 0.95, window=250)
    assert rolling.iloc[:250].isna().all()
    expected = -np.quantile(normal_returns.iloc[:250], 0.05)
    assert rolling.iloc[250] == pytest.approx(expected)


def test_kupiec_zero_breaches_is_not_zero():
    returns = pd.Series(np.full(250, 0.001))
    res = perform_kupiec_backtest(returns, 0.05, 0.99)
    assert res["actual_breaches"] == 0
    assert res["lr_stat"] == pytest.approx(-2 * 250 * np.log(0.99))


def test_kupiec_known_value():
    returns = pd.Series([-0.1] * 10 + [0.0] * 240)
    res = perform_kupiec_backtest(returns, 0.05, 0.99)
    assert res["actual_breaches"] == 10
    assert res["lr_stat"] == pytest.approx(12.955, abs=1e-3)
    assert res["test_result"].startswith("FAIL")


@pytest.mark.parametrize("breaches, zone", [(4, "GREEN"), (5, "YELLOW"), (9, "YELLOW"), (10, "RED")])
def test_basel_zones_reproduce_regulatory_table(breaches, zone):
    returns = pd.Series([-0.1] * breaches + [0.0] * (250 - breaches))
    res = perform_kupiec_backtest(returns, 0.05, 0.99)
    assert zone in res["traffic_light"]


def test_binomial_cdf_matches_scipy_for_large_samples():
    binom = pytest.importorskip("scipy.stats").binom
    assert binomial_cdf(60, 5000, 0.01) == pytest.approx(binom.cdf(60, 5000, 0.01), rel=1e-9)


@pytest.fixture
def fat_tailed_returns():
    rng = np.random.default_rng(11)
    return pd.Series(rng.standard_t(4, 3000) * 0.01)


def test_student_t_es_matches_simulation(fat_tailed_returns):
    from scipy.stats import t as student_t
    res = calculate_student_t_var(fat_tailed_returns, 1, 0.99)
    nu = res["degrees_of_freedom"]
    scale = res["sigma"] * np.sqrt((nu - 2) / nu)
    draws = res["mu"] + scale * student_t.rvs(nu, size=2_000_000, random_state=3)
    cutoff = np.quantile(draws, 0.01)
    assert res["var_daily_pct"] == pytest.approx(-cutoff, rel=0.01)
    assert res["cvar_daily_pct"] == pytest.approx(-draws[draws <= cutoff].mean(), rel=0.01)


def test_student_t_exceeds_normal_on_fat_tails(fat_tailed_returns):
    t_var = calculate_student_t_var(fat_tailed_returns, 1, 0.99)["var_daily_pct"]
    normal_var = calculate_parametric_var(fat_tailed_returns, 1, 0.99)["var_daily_pct"]
    assert t_var > normal_var


def test_cornish_fisher_tracks_fat_tails(fat_tailed_returns):
    cf = calculate_cornish_fisher_var(fat_tailed_returns, 1, 0.99)
    normal = calculate_parametric_var(fat_tailed_returns, 1, 0.99)
    assert cf["var_daily_pct"] > normal["var_daily_pct"]
    assert cf["cvar_daily_pct"] > cf["var_daily_pct"]


def test_cornish_fisher_equals_normal_without_skew_or_kurtosis(normal_returns, monkeypatch):
    monkeypatch.setattr(pd.Series, "skew", lambda self: 0.0)
    monkeypatch.setattr(pd.Series, "kurtosis", lambda self: 0.0)
    cf = calculate_cornish_fisher_var(normal_returns, 1, 0.95)["var_daily_pct"]
    normal = calculate_parametric_var(normal_returns, 1, 0.95)["var_daily_pct"]
    assert cf == pytest.approx(normal)


def test_ewma_recursion_by_hand():
    returns = pd.Series([0.01, -0.02, 0.03])
    sigma, sigma_next = ewma_volatility(returns, lam=0.94)
    v0 = np.var([0.01, -0.02, 0.03], ddof=1)
    v1 = 0.94 * v0 + 0.06 * 0.01 ** 2
    v2 = 0.94 * v1 + 0.06 * 0.02 ** 2
    v3 = 0.94 * v2 + 0.06 * 0.03 ** 2
    assert sigma.tolist() == pytest.approx(np.sqrt([v0, v1, v2]).tolist())
    assert sigma_next == pytest.approx(np.sqrt(v3))
    assert calculate_ewma_var(returns, 1, 0.99)["var_daily_pct"] == pytest.approx(2.326348 * np.sqrt(v3), rel=1e-5)


def test_ewma_reacts_to_recent_volatility():
    calm, stressed = np.full(500, 0.005), np.full(20, -0.04)
    returns = pd.Series(np.concatenate([calm * np.resize([1, -1], 500), stressed]))
    ewma = calculate_ewma_var(returns, 1, 0.99)["var_daily_pct"]
    normal = calculate_parametric_var(returns, 1, 0.99)["var_daily_pct"]
    assert ewma > normal


def test_all_models_return_positive_var_below_es(fat_tailed_returns):
    for name, res in calculate_all_var(fat_tailed_returns, 1_000_000, 0.99, num_simulations=20_000).items():
        assert 0 < res["var_daily_pct"] < res["cvar_daily_pct"], name


def test_rolling_forecasts_never_use_future_data(fat_tailed_returns):
    base = rolling_var_forecasts(fat_tailed_returns, 0.99, window=250)
    shocked_returns = fat_tailed_returns.copy()
    shocked_returns.iloc[600] = -0.5
    shocked = rolling_var_forecasts(shocked_returns, 0.99, window=250)
    pd.testing.assert_frame_equal(base.iloc[:601], shocked.iloc[:601])
    assert not np.allclose(base.iloc[601].to_numpy(), shocked.iloc[601].to_numpy())


def test_christoffersen_detects_clustered_breaches():
    clustered = np.zeros(500, dtype=int)
    clustered[100:110] = 1
    spread = np.zeros(500, dtype=int)
    spread[::50] = 1
    assert christoffersen_test(clustered)["p_value_ind"] < 0.01
    assert christoffersen_test(spread)["p_value_ind"] > 0.05


def test_christoffersen_known_value():
    hits = [0, 0, 1, 1, 0, 0, 0, 1, 0, 0]
    res = christoffersen_test(hits)
    assert (res["n00"], res["n01"], res["n10"], res["n11"]) == (4, 2, 2, 1)
    pi, pi01, pi11 = 3 / 9, 2 / 6, 1 / 3
    restricted = 6 * np.log(1 - pi) + 3 * np.log(pi)
    markov = 4 * np.log(1 - pi01) + 2 * np.log(pi01) + 2 * np.log(1 - pi11) + 1 * np.log(pi11)
    assert res["lr_ind"] == pytest.approx(max(-2 * (restricted - markov), 0.0), abs=1e-12)


def test_backtest_table_covers_every_model(fat_tailed_returns):
    forecasts = rolling_var_forecasts(fat_tailed_returns, 0.99, window=250)
    table = backtest_all_methods(fat_tailed_returns, forecasts, 0.99)
    assert list(table["Method"]) == list(forecasts.columns)
    assert (table["Test Days"] == len(fat_tailed_returns) - 250).all()
    normal = table.set_index("Method").loc["Parametric (Normal)", "Actual Breaches"]
    student = table.set_index("Method").loc["Student-t", "Actual Breaches"]
    assert student < normal


def test_beta_recovers_known_value():
    rng = np.random.default_rng(5)
    index = pd.Series(rng.normal(0, 0.01, 1000))
    stock = 1.5 * index + pd.Series(rng.normal(0, 0.002, 1000))
    assert estimate_beta(stock, index) == pytest.approx(1.5, abs=0.02)


def test_stress_scales_by_beta_and_caps_at_total_loss():
    df = run_stress_testing(1_000_000, beta=2.0).set_index("Scenario")
    assert df.loc["COVID-19 Crash (Mar 2020)", "Stock_Shock"] == pytest.approx(-0.68)
    assert df.loc["Lehman / GFC Crash (2008)", "Stock_Shock"] == -1.0
    assert df["Post_Shock_Value"].min() >= 0


def test_worst_losses_compound_returns():
    returns = pd.Series([0.0, -0.1, -0.1, 0.05])
    table = historical_worst_losses(returns, 100, horizons=(1, 2)).set_index("Horizon")
    assert table.loc["1 day", "Worst Return"] == pytest.approx(-0.1)
    assert table.loc["2 days", "Worst Return"] == pytest.approx(0.9 * 0.9 - 1)


def test_excel_report_contains_every_model_and_sheet(fat_tailed_returns):
    import io
    import openpyxl
    from excel_exporter import generate_excel_var_report

    dates = pd.bdate_range("2020-01-01", periods=len(fat_tailed_returns))
    prices = 100 * (1 + fat_tailed_returns).cumprod()
    df = pd.DataFrame({"Date": dates, "Close": prices, "Returns": prices.pct_change(),
                       "Log_Returns": np.log(prices).diff(), "Rolling_30d_Vol": prices.pct_change().rolling(30).std()})
    returns = df["Returns"].dropna()
    by_level = {cl: calculate_all_var(returns, 1_000_000, cl, 1) for cl in (0.90, 0.95, 0.99)}
    models = list(by_level[0.99])
    xlsx = generate_excel_var_report(
        "TEST", "Test Co", "USD", 1_000_000, 0.99, 1, df,
        var_by_level={m: {cl: by_level[cl][m] for cl in by_level} for m in models},
        backtest_table=backtest_all_methods(returns, rolling_var_forecasts(returns, 0.99, 250), 0.99),
        backtest_window=250, stress_df=run_stress_testing(1_000_000, 1.2),
        worst_df=historical_worst_losses(returns, 1_000_000), benchmark_name="S&P 500", beta=1.2)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Dashboard", "Backtesting", "Stress Testing", "Raw Data"]
    dash = wb["Dashboard"]
    rows = {dash.cell(row=r, column=2).value: [dash.cell(row=r, column=c).value for c in range(3, 7)] for r in range(13, 13 + len(models))}
    assert list(rows) == models
    for m in models:
        assert rows[m][0] < rows[m][1] < rows[m][2]  # 90% < 95% < 99% VaR
        assert rows[m][3] == pytest.approx(by_level[0.99][m]["cvar_scaled_amount"])
    assert wb["Raw Data"].max_row == 4 + len(df)
