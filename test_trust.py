"""Trust layer: bootstrap ranges, GARCH parameter draws, model risk, grade rules, lookbacks and the ghost effect."""

import io

import numpy as np
import pandas as pd
import pytest
from scipy.stats import norm

import trust as T
from garch import GarchParams, garch_filter, simulate_garch_paths
from var_calculator import calculate_all_var, fit_models


# ---------------------------------------------------------------
# Bootstrap
# ---------------------------------------------------------------

def test_bootstrap_coverage_on_iid_normal_returns():
    """Fast version: 150 samples of 500 i.i.d. normal days; the 90% ranges should contain the truth ~90% of the time."""
    rng = np.random.default_rng(1)
    cl, n, reps, n_boot = 0.95, 500, 150, 200
    alpha, z = 1 - cl, norm.ppf(cl)
    true_var, true_es = z, norm.pdf(z) / alpha
    hits = {"historical VaR": 0, "normal VaR": 0, "normal ES": 0}
    for _ in range(reps):
        r = rng.normal(0, 1, n)
        block, _ = T.block_length(r)
        s = r[T.stationary_bootstrap_indices(n, n_boot, block, rng)]
        hv, _ = T._empirical_rows(s, alpha)
        mu, sd = s.mean(axis=1), s.std(axis=1, ddof=1)
        for key, values, truth in (("historical VaR", hv, true_var), ("normal VaR", z * sd - mu, true_var),
                                   ("normal ES", norm.pdf(z) / alpha * sd - mu, true_es)):
            lo, hi = T._interval(values)
            hits[key] += lo <= truth <= hi
    for key, count in hits.items():
        # Binomial standard error at 150 samples is 2.4 points; allow about ±3 standard errors
        assert 0.82 <= count / reps <= 0.97, (key, count / reps)


def test_stationary_bootstrap_blocks_have_the_requested_mean_length():
    rng = np.random.default_rng(0)
    idx = T.stationary_bootstrap_indices(1000, 50, 8.0, rng)
    assert idx.min() >= 0 and idx.max() < 1000
    continues = (idx[:, 1:] == (idx[:, :-1] + 1) % 1000).mean()
    assert continues == pytest.approx(1 - 1 / 8, abs=0.01)  # a new block starts with probability 1/8
    ones = T.stationary_bootstrap_indices(1000, 50, 1.0, rng)
    assert (ones[:, 1:] == (ones[:, :-1] + 1) % 1000).mean() < 0.01  # block length 1: independent draws


def test_block_length_estimate_and_fallback(monkeypatch):
    rng = np.random.default_rng(3)
    iid = rng.normal(size=1000)
    block, method = T.block_length(iid)
    assert method == "Politis-White" and 1 <= block < 5
    persistent = np.cumsum(rng.normal(size=1000)) / 30  # strongly autocorrelated
    assert T.block_length(persistent)[0] > block

    import arch.bootstrap
    monkeypatch.setattr(arch.bootstrap, "optimal_block_length", lambda x: (_ for _ in ()).throw(RuntimeError("fail")))
    assert T.block_length(iid) == (pytest.approx(1000 ** (1 / 3)), "fallback n^(1/3)")


# ---------------------------------------------------------------
# GARCH parameter uncertainty
# ---------------------------------------------------------------

PARAMS = GarchParams(mu=0.05, omega=0.05, alpha=0.08, beta=0.88, nu=6.0)
COV = np.diag([0.02, 0.01, 0.01, 0.01, 0.8]) ** 2  # α + β = 0.96 ± 0.014: almost no draws hit the boundary


def test_parameter_draws_reproduce_the_covariance():
    draws, rejected = T.garch_parameter_draws(PARAMS, COV, 20_000, np.random.default_rng(5))
    assert len(draws) == 20_000 and rejected < 0.01
    assert draws.mean(axis=0) == pytest.approx([0.05, 0.05, 0.08, 0.88, 6.0], abs=0.02)
    assert np.cov(draws, rowvar=False).diagonal() == pytest.approx(COV.diagonal(), rel=0.08)


def test_parameter_draws_stay_valid_near_the_boundary():
    edge = GarchParams(mu=0.0, omega=0.02, alpha=0.01, beta=0.97, nu=5.0)
    draws, rejected = T.garch_parameter_draws(edge, np.diag([0.01, 0.01, 0.02, 0.02, 1.0]) ** 2, 2000, np.random.default_rng(1))
    assert rejected > 0.2
    assert (draws[:, 1] > 0).all() and (draws[:, 2] >= 0).all() and (draws[:, 2] + draws[:, 3] < 1).all()
    assert (draws[:, 4] > 2.05).all()


def test_vectorised_filter_matches_garch_filter():
    r = pd.Series(simulate_garch_paths(PARAMS, 1.0, 400, 1, seed=2)[0])
    draws = np.array([[PARAMS.mu, PARAMS.omega, PARAMS.alpha, PARAMS.beta, PARAMS.nu],
                      [0.0, 0.1, 0.05, 0.9, 8.0]])
    sigma = T.garch_next_sigma(r, draws)
    assert sigma[0] == pytest.approx(garch_filter(r, PARAMS)[-1])
    assert sigma[1] == pytest.approx(garch_filter(r, GarchParams(*draws[1]))[-1])


@pytest.fixture(scope="module")
def garch_returns():
    return pd.Series(simulate_garch_paths(PARAMS, 1.2, 1500, 1, seed=11)[0])


def test_every_model_gets_a_range(garch_returns):
    fits = fit_models(garch_returns)
    ranges = T.bootstrap_ranges(garch_returns, 0.99, fits, 2000, n_boot=150, n_boot_refit=20, n_draws=300)
    points = calculate_all_var(garch_returns, 1.0, 0.99, 1, num_simulations=2000, fits=fits)
    assert set(ranges["models"]) == set(points)
    for model, res in ranges["models"].items():
        for key in ("var", "es"):
            lo, hi = res[key]
            assert np.isfinite(lo) and 0 < lo < hi, (model, key)
    for model in ("Parametric (Normal)", "EWMA (RiskMetrics)", "GARCH(1,1)-t"):  # smooth estimators: point inside
        lo, hi = ranges["models"][model]["es"]
        assert lo <= points[model]["cvar_daily_pct"] <= hi, model
    assert ranges["models"]["GARCH(1,1)-t"]["method"] == T.METHOD_PARAMS
    assert ranges["models"]["EWMA (RiskMetrics)"]["method"] == T.METHOD_RESIDUAL


def test_garch_failure_reuses_the_ewma_range(garch_returns):
    fits = {**fit_models(garch_returns), "garch": {"converged": False, "params": None, "error": "x", "param_cov": None}}
    ranges = T.bootstrap_ranges(garch_returns, 0.95, fits, 1000, n_boot=50, n_boot_refit=5, n_draws=50)
    assert ranges["models"]["GARCH(1,1)-t"]["es"] == ranges["models"]["EWMA (RiskMetrics)"]["es"]
    assert "did not converge" in ranges["models"]["GARCH(1,1)-t"]["note"]


def test_scale_ranges_follows_each_models_horizon_rule(garch_returns):
    ten_day = calculate_all_var(garch_returns, 1_000_000, 0.99, 10, num_simulations=1000)
    # 1-day ranges at ±10% / ±20% of each model's own 1-day figure
    ranges = {"models": {m: {"var": (r["var_daily_pct"] * 0.9, r["var_daily_pct"] * 1.1),
                             "es": (r["cvar_daily_pct"] * 0.8, r["cvar_daily_pct"] * 1.2), "method": "", "note": ""}
                         for m, r in ten_day.items()}}
    table = T.scale_ranges(ranges, ten_day, 1_000_000).set_index("Model")
    for model, res in ten_day.items():
        assert table.loc[model, "VaR Low"] == pytest.approx(0.9 * res["var_scaled_amount"])
        assert table.loc[model, "ES High"] == pytest.approx(1.2 * res["cvar_scaled_amount"])


# ---------------------------------------------------------------
# Model risk
# ---------------------------------------------------------------

def _backtests(verdicts: dict, es_tests: dict = None):
    return pd.DataFrame([{"Method": m, "Verdict": v, "ES Test": (es_tests or {}).get(m, "PASS")} for m, v in verdicts.items()])


def _points(es: dict):
    return {m: {"cvar_scaled_amount": v} for m, v in es.items()}


def identity(name):
    return name


def test_model_risk_add_on_by_hand():
    points = _points({"A": 100.0, "B": 120.0, "C": 150.0, "D": 90.0})
    table = _backtests({"A": "PASS", "B": "PASS", "C": "PASS", "D": "FAIL"}, {"C": "FAIL"})
    risk = T.model_risk(points, table, "A", identity)
    assert risk["models"] == ["A", "B"]  # C fails the ES test, D the VaR tests
    assert (risk["low"], risk["high"], risk["add_on"]) == (100.0, 120.0, 20.0)
    assert risk["dispersion"] == pytest.approx(0.2)


def test_model_risk_without_a_passing_model():
    points = _points({"A": 100.0, "B": 130.0})
    risk = T.model_risk(points, _backtests({"A": "FAIL", "B": "FAIL"}), "A", identity)
    assert np.isnan(risk["add_on"]) and risk["basis"].startswith("all models")
    assert risk["dispersion"] == pytest.approx(0.3)


def test_model_risk_when_the_headline_model_fails_the_es_test():
    points = _points({"A": 100.0, "B": 130.0, "C": 110.0})
    table = _backtests({"A": "PASS", "B": "PASS", "C": "FAIL"}, {"A": "FAIL"})
    risk = T.model_risk(points, table, "A", identity)
    assert np.isnan(risk["add_on"]) and risk["basis"] == "all models (the headline model fails a backtest)"


def test_a_recommended_model_never_fails_a_backtest():
    """recommend_model and the grade agree: 'recommended' implies backtest status 'pass'; 'var_only' is graded as a fail."""
    from var_calculator import recommend_model
    table = _backtests({"A": "PASS", "B": "PASS", "C": "FAIL"}, {"A": "FAIL"}).assign(**{"Tick Loss": [0.001, 0.002, 0.0005]})
    rec = recommend_model(table)
    assert rec == {"model": "B", "status": "recommended"} and T.backtest_status(table, rec["model"]) == "pass"
    table["ES Test"] = "FAIL"
    rec = recommend_model(table)
    assert rec == {"model": "A", "status": "var_only"} and T.backtest_status(table, rec["model"]) == "fail"
    letter, reasons = T.grade(0.1, 0.1, "fail", 95, 1000, 0.0)
    assert letter == T.WORST_IF_FAILED and any("fails its backtest" in r for r in reasons)


def test_backtest_status():
    table = _backtests({"A": "PASS", "B": "FAIL", "C": "LOW POWER", "D": "PASS"}, {"D": "FAIL"})
    assert [T.backtest_status(table, m) for m in "ABCDE"] == ["pass", "fail", "low power", "fail", "not tested"]
    assert T.backtest_status(None, "A") == "not tested"


# ---------------------------------------------------------------
# Grades
# ---------------------------------------------------------------

def test_perfect_inputs_give_a_with_no_reasons():
    assert T.grade(0.10, 0.10, "pass", 95, 1000, 0.0) == ("A", [])


@pytest.mark.parametrize("kwargs, points", [
    ({"range_width": 0.20}, 0), ({"range_width": 0.21}, 1), ({"range_width": 0.40}, 1), ({"range_width": 0.41}, 2),
    ({"dispersion": 0.15}, 0), ({"dispersion": 0.30}, 1), ({"dispersion": 0.31}, 2),
    ({"backtest": "low power"}, 1), ({"backtest": "not tested"}, 1),
    ({"data_quality": 90}, 0), ({"data_quality": 89}, 1), ({"data_quality": 75}, 1), ({"data_quality": 74}, 2),
    ({"n_obs": 500}, 0), ({"n_obs": 499}, 1), ({"n_obs": 250}, 1), ({"n_obs": 249}, 2),
    ({"assumption_share": 0.25}, 0), ({"assumption_share": 0.5}, 1), ({"assumption_share": 0.51}, 2),
    ({"range_width": np.nan}, 1),
])
def test_each_rule_at_its_thresholds(kwargs, points):
    base = {"range_width": 0.1, "dispersion": 0.1, "backtest": "pass", "data_quality": 95, "n_obs": 1000, "assumption_share": 0.0}
    letter, reasons = T.grade(**{**base, **kwargs})
    assert len(reasons) == (1 if points else 0)
    if points:
        assert reasons[0].endswith(f"(−{points})")


@pytest.mark.parametrize("deductions, expected", [(1, "A"), (2, "B"), (3, "B"), (4, "C"), (5, "C"), (6, "D")])
def test_grade_bands(deductions, expected):
    # Two points per rule from range width, dispersion and sample, in that order, until the total is reached
    width = 0.5 if deductions >= 2 else 0.3 if deductions == 1 else 0.1
    disp = 0.5 if deductions >= 4 else 0.2 if deductions == 3 else 0.1
    n_obs = 100 if deductions >= 6 else 300 if deductions == 5 else 1000
    assert T.grade(width, disp, "pass", 95, n_obs, 0.0)[0] == expected


def test_failed_backtest_caps_the_grade_at_c():
    letter, reasons = T.grade(0.1, 0.1, "fail", 95, 1000, 0.0)  # 2 points alone would be B
    assert letter == "C" and any("Capped at C" in r for r in reasons)


def test_poor_data_forces_d():
    letter, reasons = T.grade(0.1, 0.1, "pass", 40, 1000, 0.0)
    assert letter == "D" and any("below 50" in r for r in reasons)


def test_model_assumptions():
    assert T.model_assumptions("Historical", 1) == (["daily returns (observed)"], [])
    inputs, assumptions = T.model_assumptions("FHS (EWMA-filtered)", 10)
    assert len(inputs) == 3 and len(assumptions) == 2  # λ and √t: two of three inputs are assumptions


def test_trusted_metric_format():
    m = T.TrustedMetric("ES", 20806.4, 18900, 23100, "B")
    assert m.format(lambda v: f"₹{v:,.0f}") == "₹20,806 (90% range ₹18,900–₹23,100, grade B)"
    assert T.TrustedMetric("x", 1.0).format() == "1 (range not available, grade not graded)"


# ---------------------------------------------------------------
# Lookback and ghost effect
# ---------------------------------------------------------------

def test_lookback_windows_have_the_expected_lengths():
    rng = np.random.default_rng(4)
    history = pd.Series(rng.normal(0, 0.01, 1300), index=pd.bdate_range("2021-01-01", periods=1300))
    table = T.lookback_sensitivity(history, 1_000_000, 0.95, 1).set_index("Lookback")
    assert table["Days"].tolist() == [252, 504, 1260, 1300] and table["Available"].all()
    assert table.loc["1y", "From"] == history.index[-252]
    assert table.loc["max", "Historical"] == pytest.approx(
        calculate_all_var(history, 1_000_000, 0.95, 1, num_simulations=1000)["Historical"]["cvar_scaled_amount"])
    assert not any(c.startswith("Monte Carlo") for c in table.columns)
    short = T.lookback_sensitivity(history.iloc[:400], 1_000_000, 0.95, 1).set_index("Lookback")
    assert short["Available"].tolist() == [True, False, False, True]


def test_lookback_flags_bad_data_and_invalid_cornish_fisher():
    from var_calculator import calculate_cornish_fisher_var
    rng = np.random.default_rng(6)
    history = pd.Series(rng.normal(0, 0.01, 1300), index=pd.bdate_range("2021-01-01", periods=1300))
    history.iloc[100], history.iloc[101] = 3.37, 1 / 4.37 - 1  # +337% then back: the Yahoo spike pattern
    full = calculate_cornish_fisher_var(history, 1_000_000, 0.95)
    assert not full["cf_valid"]
    table = T.lookback_sensitivity(history, 1_000_000, 0.95, 1).set_index("Lookback")
    assert "1 reversing spike" in table.loc["max", "Note"] and "Cornish-Fisher outside" in table.loc["max", "Note"]
    if full["cvar_scaled_amount"] <= 0:
        assert np.isnan(table.loc["max", "Cornish-Fisher"])
    else:
        assert table.loc["max", "Cornish-Fisher"] == pytest.approx(full["cvar_scaled_amount"])
    assert "reversing spike" not in table.loc["1y", "Note"] and np.isfinite(table.loc["1y", "Cornish-Fisher"])


def test_lookback_blanks_a_meaningless_cornish_fisher_es(monkeypatch):
    import var_calculator
    real = var_calculator.calculate_cornish_fisher_var

    def broken(*args, **kwargs):
        return {**real(*args, **kwargs), "cf_valid": False, "cvar_scaled_amount": -5_610_244.0}

    monkeypatch.setattr(T, "calculate_all_var", lambda *a, **k: {
        **var_calculator.calculate_all_var(*a, **k), "Cornish-Fisher": broken(a[0], a[1], a[2], a[3])})
    history = pd.Series(np.random.default_rng(1).normal(0, 0.01, 300), index=pd.bdate_range("2025-01-01", periods=300))
    table = T.lookback_sensitivity(history, 1_000_000, 0.95, 1).set_index("Lookback")
    assert np.isnan(table.loc["1y", "Cornish-Fisher"]) and "not meaningful" in table.loc["1y", "Note"]


def test_ghost_effect_finds_a_planted_loss():
    rng = np.random.default_rng(8)
    dates = pd.bdate_range("2024-01-01", periods=600)
    history = pd.Series(rng.normal(0, 0.01, 600), index=dates)
    history.iloc[300] = -0.15  # one crash
    window = history.iloc[-300:]  # the crash is the oldest day of the 300-day window
    ghost = T.ghost_effect(window, history, 0.99, horizon=21)
    assert ghost["leaving"]["Date"].tolist()[0] == dates[300]
    assert ghost["var_after"] < ghost["var_now"]
    # In the past, a crash leaves a 100-day window 100 days after it happened
    small = history.iloc[:450]
    past = T.ghost_effect(small.iloc[-100:], small, 0.99, horizon=21, jump=0.05)
    exits = past["events"][past["events"]["Cause"] == "exit"]
    assert dates[400] in set(exits["Date"]) and past["exits"] >= 1
    entry = past["events"].set_index("Date").loc[dates[300]]
    assert entry["Cause"] == "entry"


def test_excel_trust_sheet():
    import openpyxl
    from excel_exporter import generate_excel_var_report
    from var_calculator import backtest_all_methods, historical_worst_losses, rolling_var_forecasts

    rng = np.random.default_rng(2)
    returns = pd.Series(rng.standard_t(5, 600) * 0.01)
    prices = 100 * (1 + returns).cumprod()
    df = pd.DataFrame({"Date": pd.bdate_range("2023-01-02", periods=600), "Close": prices, "Returns": returns,
                       "Log_Returns": np.log1p(returns), "Rolling_30d_Vol": returns.rolling(30).std()})
    points = calculate_all_var(returns, 1_000_000, 0.95, 1, num_simulations=1000)
    ranges = {"models": {m: {"var": (r["var_daily_pct"] * 0.9, r["var_daily_pct"] * 1.1),
                             "es": (r["cvar_daily_pct"] * 0.9, r["cvar_daily_pct"] * 1.1), "method": "test", "note": ""}
                         for m, r in points.items()}}
    table = T.scale_ranges(ranges, points, 1_000_000)
    backtest = backtest_all_methods(returns, rolling_var_forecasts(returns, 0.95, 250, models=["Historical"]), 0.95)
    market = T.market_trust(points, table, backtest, "Historical", identity, 600, 92, 1, ["test"])
    history = pd.Series(returns.to_numpy(), index=df["Date"])
    xlsx = generate_excel_var_report(
        "TEST", "Test", "INR", 1_000_000, 0.95, 1, df, var_by_level={m: {0.95: r} for m, r in points.items()},
        backtest_table=backtest, backtest_window=250, stress_table=None, worst_df=historical_worst_losses(returns, 1_000_000),
        benchmark_name="Nifty 50", beta=1.0,
        trust={"ranges_table": table, "grades": market["grades"], "model_risk": market["model_risk"],
               "lookback": T.lookback_sensitivity(history, 1_000_000, 0.95, 1),
               "ghost": T.ghost_effect(history, history, 0.95, investment=1_000_000), "block_length": 1.5,
               "headline_model": "Historical"})
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Dashboard", "Backtesting", "Stress Testing", "Trust", "Raw Data"]
    dash = wb["Dashboard"]
    # Model, 95% VaR, ES, multi-day rule (columns B-E), then the five trust columns (F-J)
    headers = [dash.cell(row=15, column=c).value for c in range(2, 11)]
    assert headers[-5:] == ["95% VaR 90% Low", "95% VaR 90% High", "95% ES 90% Low", "95% ES 90% High", "Trust Grade"]
    assert dash.cell(row=16, column=6).value == pytest.approx(0.9 * points["Historical"]["var_scaled_amount"])
    assert dash.cell(row=16, column=10).value in ("A", "B", "C", "D")
    sheet = wb["Trust"]
    assert sheet.cell(row=6, column=2).value == "Historical" and sheet.cell(row=6, column=10).value in ("A", "B", "C", "D")
