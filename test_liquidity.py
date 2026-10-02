"""Liquidity pillar: hand-worked capacity and AMFI cases, Corwin-Schultz on simulated prices, costs, stressed volume,
circuit detection and the liquidity-adjusted VaR waterfall."""

import numpy as np
import pandas as pd
import pytest

import liquidity as L


def positions(rows):
    return pd.DataFrame(rows, columns=["Ticker", "Quantity", "Value"])


# ---------------------------------------------------------------
# Capacity and the AMFI test
# ---------------------------------------------------------------

THREE = positions([["A", 1000, 100_000.0], ["B", 2000, 300_000.0], ["C", 600, 100_000.0]])
THREE_ADV = {"A": 500.0, "B": 10_000.0, "C": 200.0}  # at 10%: A 20 days, B 2 days, C 30 days for the whole position


def test_amfi_pro_rata_by_hand_without_exclusion():
    res = L.amfi_stress(THREE, THREE_ADV, 0.50, participation=0.10, exclude=0.0)
    assert res["days"] == pytest.approx(15.0) and res["binding"] == "C"  # 50% of C: 300 shares at 20 a day
    assert L.amfi_stress(THREE, THREE_ADV, 0.25, 0.10, 0.0)["days"] == pytest.approx(7.5)


def test_amfi_excludes_the_least_liquid_20_percent_by_value():
    res = L.amfi_stress(THREE, THREE_ADV, 0.50, participation=0.10, exclude=0.20)
    # 20% of ₹5,00,000 = ₹1,00,000: C (the least liquid, ₹1,00,000) is removed whole; A binds at 50% × 20 days
    assert res["days"] == pytest.approx(10.0) and res["binding"] == "A"
    kept = res["table"].set_index("Ticker")["Kept Share"]
    assert kept.to_dict() == {"C": 0.0, "A": 1.0, "B": 1.0}


def test_amfi_boundary_holding_is_excluded_in_part():
    res = L.amfi_stress(THREE, THREE_ADV, 0.50, participation=0.10, exclude=0.25)
    # ₹1,25,000 removed: all of C and ₹25,000 of A, so 75% of A remains: 0.5 × 20 × 0.75 = 7.5 days
    assert res["table"].set_index("Ticker").loc["A", "Kept Share"] == pytest.approx(0.75)
    assert res["days"] == pytest.approx(7.5)


def test_holding_without_volume_never_sells():
    res = L.amfi_stress(THREE, {**THREE_ADV, "B": 0.0}, 0.5, 0.10, 0.0)
    assert np.isinf(res["days"]) and res["binding"] == "B"
    assert np.isinf(L.days_to_liquidate(100, np.nan, 0.2))


def test_sellable_share_by_hand():
    # Values 60 / 40; the first sells in 2 days, the second in 20 days
    share = L.sellable_share([60.0, 40.0], [2.0, 20.0])
    assert share[1] == pytest.approx((60 * 0.5 + 40 * 0.05) / 100)
    assert share[5] == pytest.approx((60 + 40 * 0.25) / 100)
    assert share[10] == pytest.approx((60 + 40 * 0.5) / 100)


def _frame(n=300, volume=1000.0, seed=0, start="2024-01-01"):
    rng = np.random.default_rng(seed)
    close = 100 * np.cumprod(1 + rng.normal(0, 0.01, n))
    df = pd.DataFrame({"Date": pd.bdate_range(start, periods=n), "Open": close, "High": close * 1.01,
                       "Low": close * 0.99, "Close": close, "Volume": np.full(n, volume)})
    df["Returns"] = df["Close"].pct_change()
    return df


def test_trading_capacity_uses_the_last_20_and_60_days():
    df = _frame()
    df.loc[df.index[-20:], "Volume"] = 4000.0
    cap = L.trading_capacity(positions([["X", 3000, 1e5]]).assign(Weight=1.0), {"X": df}, participation=0.2).iloc[0]
    assert cap["ADV 20d"] == 4000.0 and cap["ADV 60d"] == pytest.approx((40 * 1000 + 20 * 4000) / 60)
    assert cap["Days to Liquidate"] == pytest.approx(3000 / (0.2 * cap["ADV 60d"]))


def test_rolling_adv_uses_no_future_data():
    df = _frame().set_index("Date")
    before = L.rolling_adv(df, 63)
    changed = df.copy()
    changed.iloc[-1, changed.columns.get_loc("Volume")] = 1e9
    after = L.rolling_adv(changed, 63)
    pd.testing.assert_series_equal(before.iloc[:-1], after.iloc[:-1])


# ---------------------------------------------------------------
# Spread and costs
# ---------------------------------------------------------------

def simulated_ohlc(spread, days=500, steps=390, sigma=0.01, seed=0):
    """Efficient price as a random walk; every trade at the bid or the ask, half a spread away."""
    rng = np.random.default_rng(seed)
    high, low, close, p = [], [], [], 100.0
    for _ in range(days):
        path = p * np.exp(np.cumsum(rng.normal(0, sigma / np.sqrt(steps), steps)))
        trades = path * (1 + rng.choice([-1, 1], steps) * spread / 2)
        high.append(trades.max())
        low.append(trades.min())
        close.append(trades[-1])
        p = path[-1]
    return high, low, close


def test_corwin_schultz_recovers_a_known_spread():
    estimate = L.corwin_schultz(*simulated_ohlc(0.01))
    assert estimate.mean() == pytest.approx(0.01, abs=0.002)
    assert (estimate.dropna() >= 0).all()  # negative two-day estimates are set to zero


def test_monthly_averaging_reduces_the_upward_bias_at_tiny_spreads():
    high, low, close = simulated_ohlc(0.0, days=500, seed=1)
    df = pd.DataFrame({"Date": pd.bdate_range("2024-01-01", periods=500), "High": high, "Low": low, "Close": close})
    clipped_daily = L.corwin_schultz(high, low, close).mean()
    monthly = L.spread_profile(df, months=24)["mean"]
    assert clipped_daily > 0.003 and monthly < 0.002 and monthly < clipped_daily / 2


def test_spread_profile_uses_monthly_averages():
    high, low, close = simulated_ohlc(0.02, days=400, seed=3)
    df = pd.DataFrame({"Date": pd.bdate_range("2024-01-01", periods=400), "High": high, "Low": low, "Close": close})
    profile = L.spread_profile(df)
    assert len(profile["monthly"]) == 12 and profile["mean"] == pytest.approx(0.02, abs=0.004)
    assert np.isnan(L.spread_profile(df.drop(columns="High"))["mean"])


def test_bangia_and_square_root_impact_by_hand():
    assert L.bangia_cost(1_000_000, 0.004, 0.001, k=3) == pytest.approx(0.5 * 1_000_000 * 0.007)
    # 1 × 2% × √(10,000 / 40,000) = 1% of ₹10,00,000
    assert L.sqrt_impact(1_000_000, 0.02, 10_000, 40_000, y=1.0) == pytest.approx(10_000)
    assert L.sqrt_impact(1_000_000, 0.5, 10_000_000, 1, y=1.0) == 1_000_000  # capped at the whole value
    assert np.isnan(L.sqrt_impact(1_000_000, 0.02, 100, 0))


def test_amihud_by_hand_and_without_lookahead():
    df = _frame(n=80)
    expected = (df["Returns"].abs() / (df["Volume"] * df["Close"])).iloc[-60:].mean() * 1e7 * 1e4  # bp per ₹1 crore
    series = L.amihud(df)
    assert series.iloc[-1] == pytest.approx(expected)
    changed = df.copy()
    changed.loc[changed.index[-1], "Volume"] = 1.0
    assert L.amihud(changed).iloc[:-1].equals(series.iloc[:-1])


# ---------------------------------------------------------------
# Stressed volume
# ---------------------------------------------------------------

def test_stressed_volume_ratio_and_median():
    dates = pd.bdate_range("2019-01-01", "2021-12-31")
    volume = pd.Series(1000.0, index=dates)
    volume[(dates >= "2020-03-01") & (dates <= "2020-04-30")] = 500.0      # halves in the crisis
    volume[(dates >= "2021-06-01") & (dates <= "2021-06-30")] = 3000.0     # triples
    scenarios = pd.DataFrame({"scenario": ["Crash", "Rally", "Too early", "Too short"],
                              "start": ["2020-03-01", "2021-06-01", "2019-01-10", "2021-12-01"],
                              "end": ["2020-04-30", "2021-06-30", "2019-03-31", "2021-12-03"]})
    res = L.stressed_volume(volume, scenarios)
    table = res["table"].set_index("Scenario")
    assert table.loc["Crash", "Ratio"] == pytest.approx(0.5) and table.loc["Rally", "Ratio"] == pytest.approx(3.0)
    assert not table.loc["Too early", "Used"] and not table.loc["Too short", "Used"]
    assert res["measured"] == pytest.approx(1.75) and res["windows_used"] == 2
    assert res["factor"] == 1.0  # a crisis is never assumed to make selling easier
    crash_only = L.stressed_volume(volume, scenarios.iloc[[0]])
    assert crash_only["factor"] == pytest.approx(0.5) and crash_only["measured"] == pytest.approx(0.5)


# ---------------------------------------------------------------
# Circuit-lock risk
# ---------------------------------------------------------------

def circuit_frame(band=0.05, n=200):
    """Flat prices with three consecutive lower circuits, one isolated one and two upper circuits."""
    returns = np.zeros(n)
    for i in (50, 51, 52, 120):
        returns[i] = -band
    for i in (80, 150):
        returns[i] = band
    close = 100 * np.cumprod(1 + returns)
    df = pd.DataFrame({"Date": pd.bdate_range("2024-01-01", periods=n), "Close": close,
                       "High": close * 1.002, "Low": close * 0.998})
    for i in (50, 51, 52, 120):
        df.loc[i, "Low"] = df.loc[i, "Close"]
    for i in (80, 150):
        df.loc[i, "High"] = df.loc[i, "Close"]
    df["Returns"] = df["Close"].pct_change()
    return df


def test_band_inference_and_circuit_detection():
    df = circuit_frame()
    assert L.infer_band(df) == (0.05, 6)
    flags = L.lower_circuit_days(df, 0.05)
    assert flags.sum() == 4 and L.longest_run(flags) == 3


def test_exit_freeze_loss_compounds():
    lock = L.circuit_lock(circuit_frame())
    assert lock["source"].startswith("inferred") and lock["freeze_days"] == 3
    assert lock["loss_pct"] == pytest.approx(1 - 0.95 ** 3)
    official = L.circuit_lock(circuit_frame(), official_band=0.10, freeze_days=5)
    assert official["source"] == "official price-band file" and official["circuit_days"] == 0
    assert official["loss_pct"] == pytest.approx(1 - 0.90 ** 5)


def test_freeze_is_at_least_three_days():
    df = circuit_frame()
    df.loc[[51, 52], "Returns"] = 0.0
    lock = L.circuit_lock(df)
    assert lock["longest_run"] == 1 and lock["freeze_days"] == 3


def test_no_band_without_enough_hits():
    df = _frame()
    band, hits = L.infer_band(df)
    assert band is None and hits < L.MIN_BAND_HITS
    assert L.circuit_lock(df)["loss_pct"] == 0.0 and L.circuit_lock(df)["source"] == "no fixed band found"


def test_lvar_waterfall():
    table = L.lvar_waterfall(100.0, 10.0, 5.0, 150.0)
    assert table["Amount"].tolist() == [100.0, 10.0, 5.0, 50.0]  # circuit add-on = 150 − VaR
    assert table["Cumulative"].iloc[-1] == 165.0
    assert L.lvar_waterfall(100.0, 10.0, np.nan, 80.0)["Amount"].tolist() == [100.0, 10.0, 0.0, 0.0]


# ---------------------------------------------------------------
# The app's calculation layer
# ---------------------------------------------------------------

def test_analyse_on_synthetic_portfolio(fake_market):
    from portfolio import ENTRY_SHARES, build_positions
    from stress import load_scenarios
    from ui.liquidity_layer import analyse

    tickers = ["RELIANCE.NS", "AAPL"]
    frames = {t: fake_market(t, "2y")["df"] for t in tickers}
    long = {t: fake_market(t, "max")["df"] for t in tickers}
    pos = build_positions(pd.Series({"RELIANCE.NS": 500.0, "AAPL": 200.0}), ENTRY_SHARES,
                          {t: float(frames[t]["Close"].iloc[-1]) for t in tickers})
    res = analyse(pos, frames, long, load_scenarios(market="NIFTY50"), {}, 0.2, 0.1, 0.2, 3.0, 1.0, 0.0, 0)
    h = res["holdings"].set_index("Ticker")
    assert h.loc["AAPL", "Band Source"].startswith("not applicable") and h.loc["AAPL", "Circuit Loss"] == 0
    # The synthetic prices have no bid-ask bounce, so a zero spread estimate is correct here
    assert (h["Spread"] >= 0).all() and (h["Spread Cost"] >= 0).all() and (h["Impact Cost"] > 0).all()
    for key in ("amfi50", "sell5", "spread", "impact"):
        lo, hi = res["ranges"][key]
        assert np.isfinite(lo) and lo <= hi, key
    known = analyse(pos, frames, long, load_scenarios(market="NIFTY50"), {}, 0.2, 0.1, 0.2, 3.0, 1.0, 0.5, 0)
    assert (known["holdings"]["Spread"] == 0.005).all() and (known["holdings"]["Spread Source"] == "entered by you").all()


def test_fund_validation_reports_missing_holdings(fake_market):
    from types import SimpleNamespace
    from ui.liquidity_page import run_validation

    fund = pd.DataFrame({"ticker": ["RELIANCE.NS", "TCS.NS", "BADCO.NS"], "quantity": ["1000", "500", "10"]})
    res = run_validation(fund, SimpleNamespace(amfi_participation=0.1, amfi_exclude=0.2))
    assert res["error"] is None and res["missing"] == ["BADCO.NS"]
    assert 0 < res["days"][0.25] < res["days"][0.50]
    assert run_validation(pd.DataFrame({"Ticker": ["X"]}), SimpleNamespace())["error"].startswith("Fund portfolio")


def test_excel_liquidity_sheet(fake_market):
    import io
    import openpyxl
    from excel_exporter import generate_excel_var_report
    from portfolio import ENTRY_SHARES, build_positions
    from stress import load_scenarios
    from trust import TrustedMetric
    from ui.liquidity_layer import analyse
    from var_calculator import backtest_all_methods, calculate_all_var, historical_worst_losses, rolling_var_forecasts

    frames = {"RELIANCE.NS": fake_market("RELIANCE.NS", "2y")["df"]}
    pos = build_positions(pd.Series({"RELIANCE.NS": 500.0}), ENTRY_SHARES, {"RELIANCE.NS": float(frames["RELIANCE.NS"]["Close"].iloc[-1])})
    res = analyse(pos, frames, frames, load_scenarios(market="NIFTY50"), {}, 0.2, 0.1, 0.2, 3.0, 1.0, 0.0, 0)
    waterfall = L.lvar_waterfall(1000.0, 50.0, 20.0, 0.0)
    metrics = {"amfi50": TrustedMetric("Days to liquidate 50% (AMFI method)", 0.4, 0.3, 0.5, "A"),
               "sell5": TrustedMetric("Share sellable in 5 days", 1.0, 1.0, 1.0, "B"),
               "lvar": TrustedMetric("Liquidity-adjusted VaR", 1070.0, 900.0, 1200.0, "B", ["x (−1)"]),
               "circuit": TrustedMetric("Circuit-lock loss", 0.0, grade="C")}
    df = frames["RELIANCE.NS"]
    returns = df["Returns"].dropna()
    points = calculate_all_var(returns, 1e5, 0.95, 1, num_simulations=1000)
    xlsx = generate_excel_var_report(
        "RELIANCE.NS", "Test", "INR", 1e5, 0.95, 1, df, var_by_level={m: {0.95: r} for m, r in points.items()},
        backtest_table=backtest_all_methods(returns, rolling_var_forecasts(returns, 0.95, 250, models=["Historical"]), 0.95),
        backtest_window=250, stress_table=None, worst_df=historical_worst_losses(returns, 1e5), benchmark_name="Nifty 50",
        beta=1.0, liquidity={"metrics": metrics, "holdings": res["holdings"], "amfi": res["amfi"], "waterfall": waterfall,
                             "participation": 0.2, "bangia_k": 3.0, "impact_y": 1.0})
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert "Liquidity" in wb.sheetnames
    sheet = wb["Liquidity"]
    assert [sheet.cell(row=r, column=2).value for r in range(6, 10)] == [m.name for m in metrics.values()]
    assert sheet.cell(row=8, column=6).value == "B" and sheet.cell(row=8, column=7).value == "x (−1)"
    values = [c.value for row in sheet.iter_rows(min_col=2, max_col=4) for c in row]
    assert "+ Circuit-lock add-on" in values and "RELIANCE.NS" in values
