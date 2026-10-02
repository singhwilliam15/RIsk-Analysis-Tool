"""Credit pillar: Merton round trips, KMV, Altman by hand, point-in-time statements, ratios, flags and ratings."""

from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from scipy.stats import norm

import credit as C
import fundamentals as F

FIXTURES = Path(__file__).parent / "test_data" / "yfinance"


def saved_fund(symbol="RELIANCE.NS", info=None):
    statements = {}
    for s in F.STATEMENTS:
        frame = pd.read_csv(FIXTURES / f"{symbol}_{s}.csv", index_col=0)
        frame.columns = pd.to_datetime(frame.columns)
        statements[s] = frame
    return F.build_fundamentals(symbol, statements, info or {})


# ---------------------------------------------------------------
# Merton and KMV
# ---------------------------------------------------------------

@pytest.mark.parametrize("V, sigma_v, D, r", [(100.0, 0.25, 70.0, 0.05), (100.0, 0.40, 95.0, 0.065), (500.0, 0.15, 100.0, 0.04)])
def test_merton_recovers_known_assets(V, sigma_v, D, r):
    E, d1, _ = C.merton_equity(V, sigma_v, D, r)
    sigma_e = norm.cdf(d1) * sigma_v * V / E
    res = C.solve_merton(E, sigma_e, D, r)
    assert res["converged"]
    assert res["V"] == pytest.approx(V, rel=1e-6) and res["sigma_v"] == pytest.approx(sigma_v, rel=1e-6)
    assert res["DD"] == pytest.approx((np.log(V / D) + (r - sigma_v ** 2 / 2)) / sigma_v)
    assert res["PD"] == pytest.approx(norm.cdf(-res["DD"]))


def test_pd_rises_with_debt_and_with_volatility():
    pds_debt = [C.solve_merton(100.0, 0.4, D, 0.05)["PD"] for D in (30.0, 60.0, 120.0, 240.0)]
    pds_vol = [C.solve_merton(100.0, s, 80.0, 0.05)["PD"] for s in (0.2, 0.4, 0.6, 0.8)]
    assert all(np.diff(pds_debt) > 0) and all(np.diff(pds_vol) > 0)


def test_merton_rejects_impossible_inputs():
    assert not C.solve_merton(0.0, 0.3, 50.0, 0.05)["converged"]
    assert np.isnan(C.solve_merton(100.0, 0.3, 0.0, 0.05)["PD"])


def test_kmv_iteration_agrees_with_the_two_equation_solution():
    rng = np.random.default_rng(0)
    V0, s, D, r = 100.0, 0.25, 70.0, 0.05
    assets = V0 * np.exp(np.cumsum(rng.normal((r - s * s / 2) / 252, s / np.sqrt(252), 252)))
    equity = pd.Series([C.merton_equity(v, s, D, r)[0] for v in assets])
    kmv = C.kmv_iterative(equity, D, r)
    assert kmv["converged"] and kmv["V"] == pytest.approx(assets[-1], rel=0.01)
    assert kmv["sigma_v"] == pytest.approx(s, abs=0.025)  # one year of daily data estimates σ_V with ~1.1 pp error
    E, d1, _ = C.merton_equity(assets[-1], s, D, r)
    direct = C.solve_merton(E, norm.cdf(d1) * s * assets[-1] / E, D, r)
    assert kmv["DD"] == pytest.approx(direct["DD"], abs=0.15)


def test_default_point():
    assert C.default_point({"short_term_debt": 10.0, "long_term_debt": 40.0}) == 30.0
    assert C.default_point({"short_term_debt": 10.0, "long_term_debt": 40.0}, ltd_weight=1.0) == 50.0
    assert C.default_point({"short_term_debt": 10.0}) is None


# ---------------------------------------------------------------
# Altman
# ---------------------------------------------------------------

HAND = {"total_assets": 100.0, "current_assets": 40.0, "current_liabilities": 20.0, "retained_earnings": 30.0,
        "ebit": 10.0, "total_liabilities": 50.0, "total_equity": 50.0, "revenue": 120.0}


def test_coefficients_and_cut_offs_are_altmans():
    assert C.Z_COEFFICIENTS == (1.2, 1.4, 3.3, 0.6, 1.0) and C.Z_ZONES == (1.81, 2.99)
    assert C.Z2_COEFFICIENTS == (6.56, 3.26, 6.72, 1.05) and C.Z2_ZONES == (1.10, 2.60)


def test_altman_by_hand():
    ratios, missing = C.altman_ratios(HAND, market_cap=150.0)
    assert missing == []
    assert ratios == pytest.approx({"X1": 0.2, "X2": 0.3, "X3": 0.1, "X5": 1.2, "X4_market": 3.0, "X4_book": 1.0})
    # Z = 1.2·0.2 + 1.4·0.3 + 3.3·0.1 + 0.6·3.0 + 1.0·1.2 = 3.99; Z'' = 6.56·0.2 + 3.26·0.3 + 6.72·0.1 + 1.05·1.0 = 4.012
    assert C.altman_z(ratios) == (pytest.approx(3.99), C.SAFE)
    assert C.altman_z2(ratios) == (pytest.approx(4.012), C.SAFE)


@pytest.mark.parametrize("score, cutoffs, zone", [
    (1.80, C.Z_ZONES, C.DISTRESS), (1.81, C.Z_ZONES, C.GREY), (2.99, C.Z_ZONES, C.GREY), (3.00, C.Z_ZONES, C.SAFE),
    (1.09, C.Z2_ZONES, C.DISTRESS), (1.10, C.Z2_ZONES, C.GREY), (2.60, C.Z2_ZONES, C.GREY), (2.61, C.Z2_ZONES, C.SAFE),
])
def test_zone_boundaries(score, cutoffs, zone):
    assert C._zone(score, cutoffs) == zone


def test_never_from_partial_inputs():
    values = {k: v for k, v in HAND.items() if k != "retained_earnings"}
    ratios, missing = C.altman_ratios(values, market_cap=150.0)
    assert "retained_earnings" in missing
    assert C.altman_z(ratios) == (None, C.NOT_AVAILABLE) and C.altman_z2(ratios) == (None, C.NOT_AVAILABLE)
    # Without a market value, Z'' (book equity) still works but Z does not
    ratios, missing = C.altman_ratios(HAND, market_cap=None)
    assert "market capitalisation" in missing and C.altman_z(ratios)[0] is None and C.altman_z2(ratios)[0] is not None
    # The reported working-capital line is used when current assets/liabilities are missing
    wc_only = {**{k: v for k, v in HAND.items() if not k.startswith("current")}, "working_capital": 20.0}
    assert C.altman_ratios(wc_only, 150.0)[0]["X1"] == pytest.approx(0.2)


def test_primary_model_choice():
    config = C.load_config()
    assert C.primary_altman("Industrials", indian=False, config=config) == "Z"
    assert C.primary_altman("Technology", indian=False, config=config) == "Z''"
    assert C.primary_altman("Industrials", indian=True, config=config) == "Z''"


def test_real_statements_give_both_scores():
    fund = saved_fund()
    _, values = C.latest_year(fund)
    ratios, missing = C.altman_ratios(values, market_cap=1.6e13)
    assert missing == [] and C.altman_z(ratios)[0] is not None and C.altman_z2(ratios)[1] in (C.SAFE, C.GREY, C.DISTRESS)


# ---------------------------------------------------------------
# Point in time
# ---------------------------------------------------------------

def test_balance_sheet_is_used_only_once_public():
    fund = saved_fund()
    assert C.latest_year(fund, "2026-05-29")[0] == pd.Timestamp("2025-03-31")  # FY26 not yet public (31 Mar + 60 days)
    assert C.latest_year(fund, "2026-05-30")[0] == pd.Timestamp("2026-03-31")
    assert C.latest_year(fund, "2023-01-01") == (None, {})


def test_a_known_filing_date_makes_the_year_public_earlier():
    fund = saved_fund()
    upload = pd.DataFrame([["RELIANCE.NS", "balance", "total_assets", "2026-03-31", "1", "2026-04-25", "results filing"]],
                          columns=F.TEMPLATE_COLUMNS)
    fund, errors = F.apply_overrides(fund, upload)
    assert errors == []
    # Only total assets has a filing date; the year becomes public when its last figure is public (31 Mar + 60 days)
    assert C.latest_year(fund, "2026-04-26")[0] == pd.Timestamp("2025-03-31")
    every_field = fund["table"].assign(**{"Filing Date": pd.Timestamp("2026-04-25")})
    early = {**fund, "table": every_field}
    assert C.latest_year(early, "2026-04-25")[0] == pd.Timestamp("2026-03-31")


def test_rolling_dd_switches_balance_sheet_on_the_public_date():
    fund = saved_fund(info={"sharesOutstanding": 1.35e10})
    rng = np.random.default_rng(1)
    dates = pd.bdate_range("2025-01-01", "2026-09-30")
    close = 1300 * np.cumprod(1 + rng.normal(0, 0.012, len(dates)))
    prices = pd.DataFrame({"Date": dates, "Close": close})
    prices["Returns"] = prices["Close"].pct_change()
    rolling = C.rolling_dd(prices, fund, r=0.065)
    by_month = rolling.set_index("Date")["Balance Sheet"]
    assert (by_month[by_month.index < "2026-05-30"] <= pd.Timestamp("2025-03-31")).all()
    assert (by_month[by_month.index >= "2026-05-30"] == pd.Timestamp("2026-03-31")).all()
    assert (rolling["Shares From"] == "balance sheet").all() and rolling["DD"].notna().all()
    assert len(rolling) >= 12


# ---------------------------------------------------------------
# Ratios, flags, ratings
# ---------------------------------------------------------------

def years_frame(rows: dict) -> pd.DataFrame:
    return pd.DataFrame(rows, index=pd.to_datetime([f"{y}-03-31" for y in range(2023, 2023 + len(next(iter(rows.values()))))]))


def test_ratios_by_hand():
    years = years_frame({"total_debt": [100.0], "total_equity": [50.0], "cash": [20.0], "ebitda": [40.0], "ebit": [30.0],
                         "interest_expense": [-10.0], "current_assets": [60.0], "current_liabilities": [40.0],
                         "inventory": [20.0], "operating_cash_flow": [30.0]})
    r = C.credit_ratios(years).iloc[0]
    assert r["debt_to_equity"] == 2.0 and r["net_debt_to_ebitda"] == 2.0 and r["interest_cover"] == 3.0
    assert r["current_ratio"] == 1.5 and r["quick_ratio"] == 1.0 and r["cfo_to_ebitda"] == 0.75


def test_red_flags_on_a_weak_company():
    years = years_frame({"total_debt": [100.0, 150.0, 200.0], "total_equity": [60.0, 50.0, 40.0], "cash": [5.0, 5.0, 5.0],
                         "ebitda": [30.0, 25.0, 20.0], "ebit": [20.0, 12.0, 10.0], "interest_expense": [8.0, 9.0, 10.0],
                         "current_assets": [50.0, 45.0, 30.0], "current_liabilities": [40.0, 45.0, 50.0],
                         "inventory": [10.0, 15.0, 20.0], "operating_cash_flow": [5.0, -2.0, -4.0]})
    ratios = C.credit_ratios(years)
    assert C.negative_cfo_streak(ratios) == 2
    flags = " | ".join(C.red_flags(ratios, C.load_config()))
    for expected in ("Debt / equity: 5.00 above 2", "Net debt / EBITDA: 9.75 above 4", "Interest cover: 1.00 below 1.5",
                     "Current ratio: 0.60 below 1", "Quick ratio: 0.20 below 0.7", "Cash from operations / EBITDA",
                     "Negative cash from operations for 2 years running"):
        assert expected in flags, expected


def test_negative_equity_and_ebitda_are_flagged_as_such():
    years = years_frame({"total_debt": [50.0], "total_equity": [-10.0], "ebitda": [-5.0], "cash": [1.0]})
    flags = C.red_flags(C.credit_ratios(years), C.load_config())
    assert "Negative or zero book equity" in flags and "EBITDA zero or negative" in flags


def test_real_reliance_has_no_red_flags():
    assert C.red_flags(C.credit_ratios(C.statement_years(saved_fund())), C.load_config()) == []


def test_rating_status_as_of_a_date():
    actions = pd.DataFrame({"symbol": ["ABC", "ABC", "XYZ"], "agency": ["CRISIL", "CRISIL", "ICRA"],
                            "rating": ["CRISIL AA", "CRISIL A+", "ICRA BBB"], "action": ["reaffirmed", "downgraded", "upgraded"],
                            "action_date": pd.to_datetime(["2025-01-10", "2026-06-15", "2026-01-01"])})
    assert C.rating_status(actions, "ABC", "2026-06-14")["rating"] == "CRISIL AA"
    later = C.rating_status(actions, "ABC", "2026-06-15")
    assert later["rating"] == "CRISIL A+" and later["direction"] == "down"
    assert C.rating_status(actions, "NONE")["direction"] == "no actions on file"
    assert C.rating_status(pd.DataFrame(), "ABC")["direction"] == "not available"


def test_financial_company_detection():
    config = C.load_config()
    assert C.is_financial("Financial Services", "Banks—Regional", config)
    assert C.is_financial("Industrials", "Insurance Brokers", config)
    assert not C.is_financial("Energy", "Oil & Gas Refining & Marketing", config)


# ---------------------------------------------------------------
# The app's calculation layer
# ---------------------------------------------------------------

def test_analyse_holding_paths():
    from ui.credit_layer import analyse_holding, portfolio_view
    config = C.load_config()
    rng = np.random.default_rng(2)
    dates = pd.bdate_range("2024-10-01", "2026-09-30")
    prices = pd.DataFrame({"Date": dates, "Close": 1300 * np.cumprod(1 + rng.normal(0, 0.012, len(dates)))})
    prices["Returns"] = prices["Close"].pct_change()
    vols = {"Historical (1 year)": 0.22, "EWMA": 0.25, "GARCH(1,1)-t": 0.20}
    fund = saved_fund(info={"sharesOutstanding": 1.35e10, "financialCurrency": "INR"})
    args = (prices, fund, vols, "Energy", "Oil & Gas", "INR", pd.DataFrame(), 0.065, 0.5, 1.0, config, dates[-1])
    res = analyse_holding("RELIANCE.NS", *args)
    assert res["primary"] == "Z''" and res["z2"][0] is not None and len(res["merton"]) == 3
    assert res["merton"]["EWMA"]["PD"] > res["merton"]["GARCH(1,1)-t"]["PD"]  # more volatility, more default risk
    assert res["kmv"]["converged"] and len(res["rolling"]) > 6

    bank = analyse_holding("HDFCBANK.NS", prices, fund, vols, "Financial Services", "Banks—Regional", "INR",
                           pd.DataFrame(), 0.065, 0.5, 1.0, config, dates[-1])
    assert bank["financial"] and bank["merton"] == {} and bank["primary"] == "not applicable"

    adr = analyse_holding("RELIANCE.NS", prices, saved_fund(info={"financialCurrency": "USD", "sharesOutstanding": 1e9}),
                          vols, "Energy", "Oil & Gas", "INR", pd.DataFrame(), 0.065, 0.5, 1.0, config, dates[-1])
    assert adr["merton"] == {} and any("FX conversion" in n for n in adr["notes"])

    positions = pd.DataFrame({"Ticker": ["RELIANCE.NS", "HDFCBANK.NS"], "Value": [600.0, 400.0]})
    view = portfolio_view({"RELIANCE.NS": res, "HDFCBANK.NS": bank}, positions, "EWMA")
    pd_rel = res["merton"]["EWMA"]["PD"]
    assert view["coverage"] == pytest.approx(0.6) and view["weighted_pd"] == pytest.approx(pd_rel)
    assert view["expected_loss"] == pytest.approx(600.0 * pd_rel)  # LGD 100% for equity holders


def test_excel_credit_sheet():
    import io
    import openpyxl
    from excel_exporter import generate_excel_var_report
    from trust import TrustedMetric
    from ui.credit_layer import analyse_holding, portfolio_view
    from var_calculator import backtest_all_methods, calculate_all_var, historical_worst_losses, rolling_var_forecasts

    config = C.load_config()
    rng = np.random.default_rng(3)
    dates = pd.bdate_range("2024-10-01", "2026-09-30")
    prices = pd.DataFrame({"Date": dates, "Close": 1300 * np.cumprod(1 + rng.normal(0, 0.012, len(dates)))})
    prices["Returns"] = prices["Close"].pct_change()
    prices["Log_Returns"] = np.log1p(prices["Returns"])
    prices["Rolling_30d_Vol"] = prices["Returns"].rolling(30).std()
    vols = {"Historical (1 year)": 0.22, "EWMA": 0.25, "GARCH(1,1)-t": 0.20}
    res = analyse_holding("RELIANCE.NS", prices, saved_fund(info={"sharesOutstanding": 1.35e10}), vols, "Energy", "Oil",
                          "INR", pd.DataFrame(), 0.065, 0.5, 1.0, config, dates[-1])
    view = portfolio_view({"RELIANCE.NS": res}, pd.DataFrame({"Ticker": ["RELIANCE.NS"], "Value": [1e6]}), "EWMA")
    metrics = {"weighted_pd": TrustedMetric("Weighted PD", view["weighted_pd"], 0.001, 0.002, "B", range_label="range across x"),
               "expected_loss": TrustedMetric("Expected loss", view["expected_loss"], grade="B"),
               "weakest_dd": TrustedMetric("Lowest DD", 3.0, grade="B"), "weakest_altman": TrustedMetric("Lowest Z''", 2.2, grade="A")}
    returns = prices["Returns"].dropna()
    points = calculate_all_var(returns, 1e6, 0.95, 1, num_simulations=1000)
    xlsx = generate_excel_var_report(
        "RELIANCE.NS", "Test", "INR", 1e6, 0.95, 1, prices, var_by_level={m: {0.95: r} for m, r in points.items()},
        backtest_table=backtest_all_methods(returns, rolling_var_forecasts(returns, 0.95, 250, models=["Historical"]), 0.95),
        backtest_window=250, stress_table=None, worst_df=historical_worst_losses(returns, 1e6), benchmark_name="Nifty 50",
        beta=1.0, credit={"metrics": metrics, "view": view, "results": {"RELIANCE.NS": res}, "vol_choice": "EWMA"})
    sheet = openpyxl.load_workbook(io.BytesIO(xlsx))["Credit"]
    assert sheet.cell(row=6, column=2).value == "Weighted PD" and sheet.cell(row=6, column=6).value == "range across x"
    values = [c.value for row in sheet.iter_rows(min_col=2, max_col=8) for c in row]
    assert "RELIANCE.NS" in values and "Z''" in values


def _fake_result(ticker, dds, financial=False):
    merton = {} if financial else {name: {"DD": dd, "PD": float(norm.cdf(-dd)), "V": 1.0, "sigma_v": 0.2, "converged": True}
                                   for name, dd in zip(("Historical (1 year)", "EWMA", "GARCH(1,1)-t"), dds)}
    return {"Ticker": ticker, "merton": merton, "financial": financial, "z": (None, C.NOT_AVAILABLE),
            "z2": (2.0, C.GREY), "primary": "not applicable" if financial else "Z''", "flags": [], "notes": [],
            "rating": {"rating": None, "direction": "not available"}, "period_end": pd.Timestamp("2026-03-31"),
            "years": 4, "age_months": 6.0}


def test_tiny_pds_are_graded_on_the_distance_to_default_scale():
    from ui.credit_layer import credit_metrics, portfolio_view
    results = {"A": _fake_result("A", (11.4, 11.7, 11.2))}  # PDs around 1e-30: relative PD width would be huge
    view = portfolio_view(results, pd.DataFrame({"Ticker": ["A"], "Value": [1e6]}), "Historical (1 year)")
    metrics = credit_metrics(results, view, "Historical (1 year)", 95.0, [], C.load_config())
    pd_metric = metrics["weighted_pd"]
    assert pd_metric.range_label == "range across equity-volatility inputs"
    assert not any("range width" in r for r in pd_metric.reasons)  # (11.7 − 11.2) / 11.4 = 4%: no deduction
    assert metrics["weakest_dd"].value == pytest.approx(11.4)


def test_expected_loss_is_not_available_when_nothing_is_modelled():
    from ui.credit_layer import portfolio_view
    view = portfolio_view({"BANK": _fake_result("BANK", (), financial=True)},
                          pd.DataFrame({"Ticker": ["BANK"], "Value": [1e6]}), "EWMA")
    assert view["coverage"] == 0 and np.isnan(view["expected_loss"]) and np.isnan(view["weighted_pd"])
