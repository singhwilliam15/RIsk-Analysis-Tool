"""
Excel Exporter Module
Generates the Risk Analysis Tool's formatted multi-sheet Excel report: model comparison, portfolio risk,
out-of-sample backtest, stress tests, positions and data quality, and raw data.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np
import io
from datetime import datetime

# Styles & Palette
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")  # Dark Blue
ACCENT_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")  # Light Blue
PASS_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

BOLD_FONT = Font(name="Calibri", size=11, bold=True)
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
REGULAR_FONT = Font(name="Calibri", size=11)


def _plain(value):
    """Convert numpy / pandas scalars to plain Python values openpyxl can write."""
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        return value.item()
    return value


def _finite(value):
    """A number openpyxl can write, or 'n/a' for NaN / infinity."""
    return value if value is not None and np.isfinite(value) else "n/a"


def _write_table(ws, top_row: int, columns, rows, first_col: int = 2):
    """
    Write a header row and data rows.
    `columns` is a list of (header, number_format) pairs; '@' means text.
    Returns the next free row.
    """
    rows = list(rows)
    for c, (header, _) in enumerate(columns, start=first_col):
        cell = ws.cell(row=top_row, column=c, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r, row in enumerate(rows, start=top_row + 1):
        for c, ((_, fmt), value) in enumerate(zip(columns, row), start=first_col):
            cell = ws.cell(row=r, column=c, value=_plain(value))
            cell.font = REGULAR_FONT
            if fmt != "@":
                cell.number_format = fmt
    return top_row + len(rows) + 2


def generate_excel_var_report(symbol: str, company_name: str, currency: str, investment: float,
                              confidence_level: float, holding_period: int, df_data: pd.DataFrame,
                              var_by_level: dict, backtest_table: pd.DataFrame, backtest_window: int,
                              stress_table: pd.DataFrame, worst_df: pd.DataFrame,
                              benchmark_name: str, beta: float, portfolio: dict = None,
                              recommendation: dict = None, data_note: str = None, risk_free_rate: float = None,
                              vol_shock_table: pd.DataFrame = None, beta_down: float = float("nan"),
                              data_layer: dict = None, trust: dict = None, liquidity: dict = None,
                              credit: dict = None, concentration: dict = None) -> bytes:
    """
    Generate the Risk Analysis Tool's Excel report and return it as .xlsx bytes.
    `var_by_level` maps model name -> {confidence level -> VaR result dict}.
    `stress_table` comes from stress.run_scenarios (None if the market history was unavailable).
    `portfolio` (optional) holds "decomposition" (portfolio.risk_decomposition), "diversification",
    "correlation" and "alignment" (portfolio.alignment_report).
    `data_layer` (optional) holds "positions" (portfolio.build_positions), "quality" (ticker ->
    data_quality.assess_holding result), "volume_sources" and "prices_as_of"; it adds a Positions & Data sheet.
    `trust` (optional) holds "ranges_table" (trust.scale_ranges), "grades", "model_risk", "lookback", "ghost",
    "block_length" and "headline_model"; it adds 90% ranges and grades to the Dashboard and a Trust sheet.
    `liquidity` (optional) holds "metrics" (TrustedMetric per headline), "holdings", "amfi", "waterfall",
    "participation", "bangia_k" and "impact_y"; it adds a Liquidity sheet.
    `credit` (optional) holds "metrics", "view" (credit_layer.portfolio_view), "results" and "vol_choice"; it adds a
    Credit sheet. `concentration` (optional) holds "result" (concentration_layer.analyse), "metrics" and
    "factor_meta"; it adds a Concentration sheet.
    """
    money = f'"{currency}" #,##0'
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # -------------------------------------------------------------
    # SHEET 1: DASHBOARD
    # -------------------------------------------------------------
    ws_dash = wb.create_sheet(title="Dashboard")
    ws_dash["B2"] = f"⚡ RISK ANALYSIS TOOL — MARKET RISK DASHBOARD ({symbol})"
    ws_dash["B2"].font = TITLE_FONT
    ws_dash["B3"] = f"Company: {company_name} | Currency: {currency} | Generated: {datetime.now():%Y-%m-%d %H:%M}"
    ws_dash["B3"].font = SUBTITLE_FONT

    ws_dash["B5"] = "POSITION SETTINGS"
    ws_dash["B5"].font = BOLD_FONT
    ws_dash["B5"].fill = ACCENT_FILL
    settings_data = [
        ("Position Value", investment, money),
        ("Selected Confidence Level", confidence_level, "0.0%"),
        ("Holding Period (Days)", holding_period, "0"),
        (f"Beta vs {benchmark_name}", beta if np.isfinite(beta) else "n/a", "0.00"),
        ("Downside Beta (worst 10% of days)", beta_down if np.isfinite(beta_down) else "n/a", "0.00"),
        ("Risk-free Rate (assumption)", risk_free_rate if risk_free_rate is not None else "n/a", "0.00%"),
        ("Data Source", data_note or "n/a", "@"),
    ]
    for i, (label, val, fmt) in enumerate(settings_data, start=6):
        ws_dash[f"B{i}"] = label
        ws_dash[f"C{i}"] = _plain(val)
        ws_dash[f"C{i}"].number_format = fmt
        ws_dash[f"C{i}"].font = BOLD_FONT

    ws_dash["B14"] = f"VAR AND EXPECTED SHORTFALL BY MODEL ({holding_period}-DAY HORIZON)"
    ws_dash["B14"].font = BOLD_FONT
    cl_pct = f"{confidence_level * 100:g}%"
    levels_shown = sorted(next(iter(var_by_level.values())))
    trust_columns, trust_rows = [], {}
    if trust is not None:
        # "value (90% range, grade)" for every headline number: the range and grade sit next to each model's figures
        trust_columns = [(f"{cl_pct} VaR 90% Low", money), (f"{cl_pct} VaR 90% High", money),
                         (f"{cl_pct} ES 90% Low", money), (f"{cl_pct} ES 90% High", money), ("Trust Grade", "@")]
        grades = trust["grades"].set_index("Model")["Grade"]
        for row in trust["ranges_table"].to_dict("records"):
            trust_rows[row["Model"]] = [_finite(row["VaR Low"]), _finite(row["VaR High"]), _finite(row["ES Low"]),
                                        _finite(row["ES High"]), grades.get(row["Model"], "n/a")]
    _write_table(
        ws_dash, 15,
        [("Model", "@")] + [(f"{cl * 100:g}% VaR", money) for cl in levels_shown]
        + [(f"{cl_pct} Expected Shortfall", money), ("Multi-day Rule", "@")] + trust_columns,
        [
            (model, *[levels[cl]["var_scaled_amount"] for cl in levels_shown], levels[confidence_level]["cvar_scaled_amount"],
             levels[confidence_level]["scaling_rule"] if holding_period > 1 else "1 day",
             *trust_rows.get(model, ["n/a"] * len(trust_columns)))
            for model, levels in var_by_level.items()
        ],
    )

    # -------------------------------------------------------------
    # PORTFOLIO SHEET (portfolio mode only)
    # -------------------------------------------------------------
    if portfolio is not None:
        ws_port = wb.create_sheet(title="Portfolio Risk")
        decomposition = portfolio["decomposition"]
        basis = decomposition["basis"]
        ws_port["B2"] = f"🧩 PORTFOLIO RISK DECOMPOSITION — {basis.upper()} ({cl_pct}, {holding_period}-DAY)"
        ws_port["B2"].font = TITLE_FONT
        ws_port["B3"] = (f"Component {basis} is the Euler allocation and sums exactly to the portfolio total. "
                         "Incremental = portfolio risk with the holding minus without it.")
        ws_port["B3"].font = SUBTITLE_FONT

        div = portfolio["diversification"]
        alignment = portfolio.get("alignment") or {}
        summary_rows = [
            (f"Portfolio {basis}", decomposition["total"], money),
            ("Sum of Standalone", decomposition["standalone_sum"], money),
            ("Diversification Benefit", decomposition["diversification_benefit"], money),
            ("Average Pairwise Correlation", div["average_correlation"], "0.00"),
            ("Common Trading Days", div["common_days"], "0"),
            ("Holding Limiting the Sample", alignment.get("limiting_ticker", "n/a"), "@"),
        ]
        for i, (label, val, fmt) in enumerate(summary_rows, start=5):
            ws_port[f"B{i}"] = label
            ws_port[f"C{i}"] = _plain(val)
            ws_port[f"C{i}"].number_format = fmt
            ws_port[f"C{i}"].font = BOLD_FONT

        comp = decomposition["table"]
        comp_cols = [
            ("Ticker", "@"), ("Company", "@"), ("Weight", "0.0%"), ("Annualized Volatility", "0.0%"),
            ("Standalone", money), ("Component", money), ("Contribution %", "0.0%"),
            ("Risk / Weight", "0.00"), ("Incremental", money),
        ]
        next_row = _write_table(ws_port, 12, comp_cols, comp[[name for name, _ in comp_cols]].itertuples(index=False))
        total_row = next_row - 1
        ws_port.cell(row=total_row, column=2, value="TOTAL").font = BOLD_FONT
        for name in ("Weight", "Component", "Contribution %"):
            col = 2 + [n for n, _ in comp_cols].index(name)
            letter = get_column_letter(col)
            cell = ws_port.cell(row=total_row, column=col, value=f"=SUM({letter}13:{letter}{12 + len(comp)})")
            cell.number_format = dict(comp_cols)[name]
            cell.font = BOLD_FONT

        ws_port.cell(row=total_row + 2, column=2, value="CORRELATION OF DAILY RETURNS").font = BOLD_FONT
        corr = portfolio["correlation"]
        _write_table(
            ws_port, total_row + 3,
            [("", "@")] + [(t, "0.00") for t in corr.columns],
            [(t, *row) for t, row in zip(corr.index, corr.to_numpy())],
        )

    # -------------------------------------------------------------
    # SHEET 2: BACKTESTING
    # -------------------------------------------------------------
    ws_back = wb.create_sheet(title="Backtesting")
    ws_back["B2"] = f"🔬 OUT-OF-SAMPLE BACKTEST — {cl_pct} VAR, ROLLING {backtest_window}-DAY WINDOW"
    ws_back["B2"].font = TITLE_FONT
    ws_back["B3"] = ("Kupiec: correct breach count. Christoffersen: breaches independent. "
                     "Conditional coverage: both. PASS requires every p-value >= 0.05. "
                     "Recommended = lowest tick loss among PASS models. "
                     "ES test (McNeil-Frey): breach-day (loss - ES)/sigma should average zero; low p = ES too small.")
    ws_back["B3"].font = SUBTITLE_FONT
    rec = recommendation or {"model": None, "status": "low_power"}
    ws_back["B4"] = {
        "recommended": f"Recommended model: {rec['model']}",
        "none_pass": f"No model passes all tests; lowest tick loss: {rec['model']} (use with caution)",
    }.get(rec["status"], "Not enough out-of-sample data for a meaningful backtest; choose a longer lookback.")
    ws_back["B4"].font = BOLD_FONT
    if backtest_table is not None and len(backtest_table):
        backtest_cols = [
            ("Method", "@"), ("Test Days", "0"), ("Expected Breaches", "0.0"), ("Actual Breaches", "0"),
            ("Verdict", "@"), ("Tick Loss", "0.000000"), ("Breach Rate", "0.00%"), ("Kupiec p-value", "0.000"),
            ("Independence p-value", "0.000"), ("Conditional Coverage p-value", "0.000"),
            ("Back-to-Back Breaches", "0"), ("Traffic Light", "@"), ("Avg VaR", "0.00%"),
        ]
        if "ES Test" in backtest_table.columns:
            backtest_cols += [("ES Test", "@"), ("ES p-value", "0.000"), ("ES Mean Residual", "0.000")]
        _write_table(ws_back, 6, backtest_cols, backtest_table[[name for name, _ in backtest_cols]].itertuples(index=False))
        verdict_col = 2 + [name for name, _ in backtest_cols].index("Verdict")
        for r in range(7, 7 + len(backtest_table)):
            cell = ws_back.cell(row=r, column=verdict_col)
            if cell.value in ("PASS", "FAIL"):
                cell.fill = PASS_FILL if cell.value == "PASS" else FAIL_FILL
            cell.font = BOLD_FONT

    # -------------------------------------------------------------
    # SHEET 3: STRESS TESTING
    # -------------------------------------------------------------
    ws_stress = wb.create_sheet(title="Stress Testing")
    ws_stress["B2"] = "⚡ STRESS TESTING — HISTORICAL CRISIS SCENARIOS"
    ws_stress["B2"].font = TITLE_FONT
    ws_stress["B3"] = (f"Market falls are measured peak to trough from downloaded {benchmark_name} prices. Historical replay = the "
                       "position's actual return between the same dates; beta-proxy = market fall x downside beta, used only "
                       "when the position has no prices for the period.")
    ws_stress["B3"].font = SUBTITLE_FONT
    next_row = 5
    if stress_table is not None and "Peak" in stress_table.columns:
        covered = stress_table[stress_table["Peak"].notna()]
        next_row = _write_table(
            ws_stress, 5,
            [("Scenario", "@"), ("Description", "@"), ("Peak", "@"), ("Trough", "@"), ("Market Drawdown", "0.0%"),
             ("Market Recovery (Days)", "@"), ("Method", "@"), ("Position Return", "0.0%"), ("P&L", money),
             ("Post-Shock Value", money)],
            [(r["Scenario"], r["Description"], f"{r['Peak']:%Y-%m-%d}", f"{r['Trough']:%Y-%m-%d}", r["Market Drawdown"],
              "not yet" if pd.isna(r["Market Recovery (days)"]) else int(r["Market Recovery (days)"]),
              r["Method"], r["Position Return"], r["P&L"], r["Post-Shock Value"]) for _, r in covered.iterrows()],
        )
    if vol_shock_table is not None:
        ws_stress.cell(row=next_row, column=2, value=f"VOLATILITY SHOCK ({cl_pct}, {holding_period}-DAY)").font = BOLD_FONT
        next_row = _write_table(
            ws_stress, next_row + 1,
            [("Volatility", "@"), ("Historical VaR", money), ("Historical ES", money), ("Normal VaR", money), ("Normal ES", money)],
            vol_shock_table[["Volatility", "Historical VaR", "Historical ES", "Normal VaR", "Normal ES"]].itertuples(index=False),
        )
    ws_stress.cell(row=next_row, column=2, value="WORST ACTUAL LOSSES IN SAMPLE").font = BOLD_FONT
    window_end = df_data.loc[worst_df["Window End"], "Date"].dt.strftime("%Y-%m-%d").to_numpy() if len(worst_df) else []
    _write_table(
        ws_stress, next_row + 1,
        [("Horizon", "@"), ("Worst Return", "0.00%"), ("Loss", money), ("Window End", "@")],
        zip(worst_df["Horizon"], worst_df["Worst Return"], worst_df["Loss"], window_end),
    )

    # -------------------------------------------------------------
    # POSITIONS AND DATA QUALITY (when the app supplies them)
    # -------------------------------------------------------------
    if data_layer is not None:
        ws_data = wb.create_sheet(title="Positions & Data")
        ws_data["B2"] = "💼 POSITIONS AND DATA QUALITY"
        ws_data["B2"].font = TITLE_FONT
        ws_data["B3"] = (f"Positions converted at the latest close ({data_layer['prices_as_of']:%Y-%m-%d}). Data-quality "
                         "score = 100 minus penalties; the thresholds are assumptions (docs/methodology.md, section 7).")
        ws_data["B3"].font = SUBTITLE_FONT
        positions = data_layer["positions"]
        next_row = _write_table(
            ws_data, 5,
            [("Ticker", "@"), ("Quantity", "#,##0.00"), ("Price", "#,##0.00"), ("Value", money), ("Weight", "0.0%"),
             ("Sector", "@")],
            positions[["Ticker", "Quantity", "Price", "Value", "Weight", "Sector"]].itertuples(index=False),
        )
        ws_data.cell(row=next_row, column=2, value="DATA QUALITY BY HOLDING").font = BOLD_FONT
        quality = data_layer["quality"]
        volume_sources = data_layer.get("volume_sources", {})
        _write_table(
            ws_data, next_row + 1,
            [("Ticker", "@"), ("Score (0-100)", "0"), ("Volume From", "@"), ("Issues", "@")],
            [(t, q["score"], " + ".join(volume_sources.get(t, [t])), "; ".join(q["reasons"]) or "no issues found")
             for t, q in quality.items()],
        )

    # -------------------------------------------------------------
    # TRUST (when the app supplies it)
    # -------------------------------------------------------------
    if trust is not None:
        ws_trust = wb.create_sheet(title="Trust")
        ws_trust["B2"] = "🛡️ TRUST: 90% RANGES, GRADES AND MODEL RISK"
        ws_trust["B2"].font = TITLE_FONT
        ws_trust["B3"] = (f"{cl_pct}, {holding_period}-day. Ranges: 5th-95th percentile of block-bootstrap resamples "
                          f"(mean block {trust['block_length']:.1f} days) or GARCH parameter draws. Grades follow the "
                          "rules in docs/methodology.md section 8.")
        ws_trust["B3"].font = SUBTITLE_FONT
        grades = trust["grades"].set_index("Model")
        next_row = _write_table(
            ws_trust, 5,
            [("Model", "@"), ("VaR", money), ("VaR 90% Low", money), ("VaR 90% High", money), ("ES", money),
             ("ES 90% Low", money), ("ES 90% High", money), ("Backtest", "@"), ("Grade", "@"), ("Range Method", "@"),
             ("Reasons", "@")],
            [(r["Model"], r["VaR"], _finite(r["VaR Low"]), _finite(r["VaR High"]), r["ES"], _finite(r["ES Low"]),
              _finite(r["ES High"]), grades.loc[r["Model"], "Backtest"], grades.loc[r["Model"], "Grade"], r["Range Method"],
              grades.loc[r["Model"], "Reasons"]) for r in trust["ranges_table"].to_dict("records")],
        )
        risk = trust["model_risk"]
        ws_trust.cell(row=next_row, column=2, value="MODEL RISK").font = BOLD_FONT
        next_row = _write_table(
            ws_trust, next_row + 1, [("Measure", "@"), ("Value", money)],
            [(f"Lowest ES across {risk['basis']}", risk["low"]), (f"Highest ES across {risk['basis']}", risk["high"]),
             ("Recommended model's ES", _finite(risk["recommended_es"])),
             ("Model-risk add-on (highest passing ES − recommended ES)", _finite(risk["add_on"]))],
        )
        lookback = trust["lookback"]
        models = [c for c in lookback.columns if c not in ("Lookback", "Days", "Available", "From", "Note")]
        ws_trust.cell(row=next_row, column=2, value="LOOKBACK SENSITIVITY (ES)").font = BOLD_FONT
        next_row = _write_table(
            ws_trust, next_row + 1, [("Lookback", "@"), ("Days", "0")] + [(m, money) for m in models] + [("Note", "@")],
            [(r["Lookback"], r["Days"], *[_finite(r.get(m, np.nan)) for m in models], r.get("Note") or "")
             for r in lookback.to_dict("records")],
        )
        ghost = trust["ghost"]
        ws_trust.cell(row=next_row, column=2, value="GHOST EFFECT IN HISTORICAL VAR").font = BOLD_FONT
        _write_table(
            ws_trust, next_row + 1, [("Measure", "@"), ("Value", "@")],
            [("Window (returns)", ghost["window_days"]), ("Historical VaR now", round(ghost["var_now"])),
             (f"Without the oldest {ghost['horizon']} days", _finite(round(ghost["var_after"])) if np.isfinite(ghost["var_after"]) else "n/a"),
             (f"Tail losses leaving within {ghost['horizon']} days", len(ghost["leaving"])),
             (f"Past VaR jumps > {ghost['jump']:.0%} caused by a loss leaving", ghost["exits"])],
        )

    # -------------------------------------------------------------
    # LIQUIDITY (when the app supplies it)
    # -------------------------------------------------------------
    if liquidity is not None:
        ws_liq = wb.create_sheet(title="Liquidity")
        ws_liq["B2"] = "💧 LIQUIDITY RISK"
        ws_liq["B2"].font = TITLE_FONT
        ws_liq["B3"] = (f"Participation {liquidity['participation']:.0%} of daily volume, Bangia k {liquidity['bangia_k']:g}, "
                        f"impact constant Y {liquidity['impact_y']:g}: assumptions. Methods in docs/methodology.md section 9.")
        ws_liq["B3"].font = SUBTITLE_FONT
        metrics = liquidity["metrics"]
        units = {"amfi50": "0.0", "sell5": "0%", "lvar": money, "circuit": money}
        next_row = _write_table(
            ws_liq, 5, [("Headline", "@"), ("Value", "@"), ("90% Low", "@"), ("90% High", "@"), ("Grade", "@"), ("Reasons", "@")],
            [(m.name, _finite(m.value), _finite(m.low), _finite(m.high), m.grade, "; ".join(m.reasons) or "no deductions")
             for m in metrics.values()],
        )
        for r, key in enumerate(metrics, start=6):
            for c in (3, 4, 5):
                ws_liq.cell(row=r, column=c).number_format = units[key]
        h = liquidity["holdings"]
        ws_liq.cell(row=next_row, column=2, value="BY HOLDING").font = BOLD_FONT
        next_row = _write_table(
            ws_liq, next_row + 1,
            [("Ticker", "@"), ("Value", money), ("ADV 60d", "#,##0"), ("Days to Liquidate", "0.0"), ("Stress Factor", "0.00"),
             ("Stressed Days", "0.0"), ("Spread", "0.00%"), ("Spread Source", "@"), ("Spread Cost", money),
             ("Impact Cost", money), ("Band", "0%"), ("Band Source", "@"), ("Lower-Circuit Days", "0"),
             ("Freeze Days", "0"), ("Circuit Loss", money)],
            [(r["Ticker"], r["Value"], _finite(r["ADV 60d"]), _finite(r["Days to Liquidate"]), _finite(r["Stress Factor"]),
              _finite(r["Stressed Days"]), _finite(r["Spread"]), r["Spread Source"], _finite(r["Spread Cost"]),
              _finite(r["Impact Cost"]), _finite(r["Band"]), r["Band Source"], r["Lower-Circuit Days"], r["Freeze Days"],
              r["Circuit Loss"]) for r in h.to_dict("records")],
        )
        ws_liq.cell(row=next_row, column=2, value="SEBI/AMFI-STYLE STRESS TEST (DAYS)").font = BOLD_FONT
        amfi = liquidity["amfi"]
        next_row = _write_table(
            ws_liq, next_row + 1, [("Portfolio Sold", "@"), ("Least Liquid Excluded", "0.0"), ("Nothing Excluded", "0.0")],
            [(f"{f:.0%}", _finite(amfi[(f, True)]["days"]), _finite(amfi[(f, False)]["days"])) for f in (0.25, 0.50)],
        )
        ws_liq.cell(row=next_row, column=2, value="LIQUIDITY-ADJUSTED VAR").font = BOLD_FONT
        _write_table(ws_liq, next_row + 1, [("Step", "@"), ("Amount", money), ("Cumulative", money)],
                     liquidity["waterfall"][["Step", "Amount", "Cumulative"]].itertuples(index=False))

    # -------------------------------------------------------------
    # CREDIT (when the app supplies it)
    # -------------------------------------------------------------
    if credit is not None:
        ws_cr = wb.create_sheet(title="Credit")
        ws_cr["B2"] = "🏦 CREDIT RISK"
        ws_cr["B2"].font = TITLE_FONT
        ws_cr["B3"] = ("PD is a model-implied, risk-neutral Merton probability, not an agency PD. Expected loss assumes "
                       f"loss given default of 100% for equity holders. Equity volatility: {credit['vol_choice']}. "
                       "Methods in docs/methodology.md section 10.")
        ws_cr["B3"].font = SUBTITLE_FONT
        metrics = credit["metrics"]
        next_row = _write_table(
            ws_cr, 5, [("Headline", "@"), ("Value", "0.0000"), ("Range Low", "0.0000"), ("Range High", "0.0000"),
                       ("Range", "@"), ("Grade", "@"), ("Reasons", "@")],
            [(m.name, _finite(m.value), _finite(m.low), _finite(m.high), m.range_label, m.grade,
              "; ".join(m.reasons) or "no deductions") for m in metrics.values()],
        )
        formats = {"weighted_pd": "0.000%", "expected_loss": money, "weakest_dd": "0.00", "weakest_altman": "0.00"}
        for r, key in enumerate(metrics, start=6):
            for c in (3, 4, 5):
                ws_cr.cell(row=r, column=c).number_format = formats.get(key, "0.00")
        t = credit["view"]["table"]
        ws_cr.cell(row=next_row, column=2, value="BY HOLDING").font = BOLD_FONT
        next_row = _write_table(
            ws_cr, next_row + 1,
            [("Ticker", "@"), ("Value", money), ("DD", "0.00"), ("PD", "0.000%"), ("Altman Model", "@"), ("Altman Score", "0.00"),
             ("Altman Zone", "@"), ("Red Flags", "0"), ("Rating", "@"), ("Rating Direction", "@"), ("Balance Sheet (FY end)", "@")],
            [(r["Ticker"], r["Value"], _finite(r["DD"]), _finite(r["PD"]), r["Altman Model"], _finite(r["Altman Score"]),
              r["Altman Zone"], r["Red Flags"], r["Rating"], r["Rating Direction"],
              f"{r['Balance Sheet']:%Y-%m-%d}" if pd.notna(r["Balance Sheet"]) else "not available")
             for r in t.to_dict("records")],
        )
        flag_rows = [(tk, f) for tk, res in credit["results"].items() for f in res["flags"] + res["notes"]]
        if flag_rows:
            ws_cr.cell(row=next_row, column=2, value="RED FLAGS AND NOTES").font = BOLD_FONT
            _write_table(ws_cr, next_row + 1, [("Ticker", "@"), ("Flag or note", "@")], flag_rows)

    # -------------------------------------------------------------
    # CONCENTRATION AND FACTORS (when the app supplies it)
    # -------------------------------------------------------------
    if concentration is not None:
        ws_cf = wb.create_sheet(title="Concentration")
        res, meta = concentration["result"], concentration.get("factor_meta") or {}
        ws_cf["B2"] = "🧭 CONCENTRATION AND FACTOR RISK"
        ws_cf["B2"].font = TITLE_FONT
        ws_cf["B3"] = (f"Model: {res['model']}. Factor data: {meta.get('source', 'not available')} to "
                       f"{meta.get('end_date', 'n/a')}. Methods in docs/methodology.md section 11.")
        ws_cf["B3"].font = SUBTITLE_FONT
        next_row = _write_table(
            ws_cf, 5, [("Headline", "@"), ("Value", "0.00"), ("Low", "0.00"), ("High", "0.00"), ("Range", "@"),
                       ("Grade", "@"), ("Reasons", "@")],
            [(m.name, _finite(m.value), _finite(m.low), _finite(m.high), m.range_label, m.grade,
              "; ".join(m.reasons) or "no deductions") for m in concentration["metrics"].values()],
        )
        fits = {t: f for t, f in res["fits"].items() if f is not None}
        if fits:
            names = list(next(iter(fits.values()))["betas"].index)
            ws_cf.cell(row=next_row, column=2, value="FACTOR EXPOSURES (NEWEY-WEST t IN BRACKETS)").font = BOLD_FONT
            next_row = _write_table(
                ws_cf, next_row + 1, [("Ticker", "@"), ("Days", "0")] + [(n, "@") for n in names]
                + [("Alpha (annual)", "0.0%"), ("R²", "0.00")],
                [(t, f["nobs"], *[f"{f['betas'][n]:+.2f} ({f['t_nw'][n]:+.1f})" for n in names], f["alpha_annual"], f["r2"])
                 for t, f in fits.items()],
            )
        fr = res["factor_risk"]
        if fr is not None:
            ws_cf.cell(row=next_row, column=2, value="EULER SPLIT OF PORTFOLIO VARIANCE").font = BOLD_FONT
            next_row = _write_table(ws_cf, next_row + 1, [("Source", "@"), ("Share", "0.0%"), ("Parametric VaR Part", money)],
                                    [(k, fr["shares"][k], fr["var_parts"][k]) for k in fr["shares"].index])
        if "sectors" in res:
            ws_cf.cell(row=next_row, column=2, value=f"SECTORS (HHI {res['hhi']:.3f}, EFFECTIVE BETS {res['pca']['enb']:.2f})").font = BOLD_FONT
            next_row = _write_table(ws_cf, next_row + 1, [("Sector", "@"), ("Weight", "0.0%"), ("Risk Share", "0.0%")],
                                    res["sectors"][["Sector", "Weight", "Risk Share"]].itertuples(index=False))
        if res.get("crisis") is not None:
            t = res["crisis"]["table"]
            ws_cf.cell(row=next_row, column=2, value="CORRELATION IN A CRISIS").font = BOLD_FONT
            _write_table(ws_cf, next_row + 1, [("Sample", "@"), ("Days", "0"), ("Average Correlation", "0.00"),
                                               ("Portfolio ES", money), ("Benefit Kept", "0%")],
                         [(r["Sample"], r["Days"], _finite(r["Average Correlation"]), _finite(r["ES"]), _finite(r["Benefit Kept"]))
                          for r in t.to_dict("records")])

    # -------------------------------------------------------------
    # SHEET 4: RAW DATA
    # -------------------------------------------------------------
    ws_raw = wb.create_sheet(title="Raw Data")
    ws_raw["B2"] = f"📈 HISTORICAL MARKET PRICE DATA ({symbol})"
    ws_raw["B2"].font = TITLE_FONT
    raw = df_data[["Date", "Close", "Returns", "Log_Returns", "Rolling_30d_Vol"]].copy()
    raw["Date"] = raw["Date"].dt.strftime("%Y-%m-%d")
    raw = raw.astype(object).where(raw.notna(), None)
    _write_table(
        ws_raw, 4,
        [("Date", "@"), ("Close Price", "#,##0.00"), ("Daily Return", "0.00%"), ("Log Return", "0.00%"),
         ("Rolling 30d Volatility", "0.00%")],
        raw.itertuples(index=False),
    )

    # Fit column widths to content, ignoring the long title and note rows
    for ws in wb.worksheets:
        for col in ws.iter_cols(min_row=4):
            max_len = max((len(str(cell.value)) for cell in col if cell.value is not None), default=0)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 3, 12), 45)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
